"""Acrescenta ao Excel uma visão navegável dos alunos agrupados por STATUS_CURSO."""
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from referencia import CATEGORIAS_REFERENCIA

COLUNAS = [
    "STATUS_CURSO", "Aluno", "Identificador", "Curso", "Status Atual",
    "Período de coleta", "Página",
]
LARGURAS = {
    "STATUS_CURSO": 28,
    "Aluno": 52,
    "Identificador": 18,
    "Curso": 38,
    "Status Atual": 26,
    "Período de coleta": 22,
    "Página": 10,
}


def adicionar_visao_por_categoria(caminho_excel, resultado):
    """Cria a aba 'Por Categoria' e liga a aba Referência aos primeiros alunos de cada grupo."""
    auditoria = resultado.get("classificacao", {}).get("auditoria") or []
    registros = resultado.get("registros") or []
    if len(auditoria) != len(registros):
        raise ValueError("Não foi possível montar a visão por categoria: classificação e registros estão desalinhados.")

    wb = load_workbook(caminho_excel)
    if "Por Categoria" in wb.sheetnames:
        del wb["Por Categoria"]
    ws = wb.create_sheet("Por Categoria")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(COLUNAS)

    ordem = {categoria: i for i, categoria in enumerate(CATEGORIAS_REFERENCIA)}
    linhas = []
    for registro, classificacao in zip(registros, auditoria):
        categoria = classificacao.get("categoria_final") or "NÃO LOCALIZADO"
        linhas.append((
            categoria,
            registro.get("Aluno", ""),
            registro.get("Identificador", ""),
            registro.get("Curso", ""),
            registro.get("Status Atual", ""),
            registro.get("Período de coleta", ""),
            registro.get("Página", ""),
        ))

    linhas.sort(key=lambda row: (ordem.get(row[0], 999), str(row[1]).casefold(), str(row[2])))
    primeira_linha = {}
    for row in linhas:
        ws.append(row)
        primeira_linha.setdefault(row[0], ws.max_row)

    header_fill = PatternFill("solid", fgColor="243D55")
    alt_fill = PatternFill("solid", fgColor="F0F4F7")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 30

    for row_index in range(2, ws.max_row + 1):
        if row_index % 2 == 0:
            for cell in ws[row_index]:
                cell.fill = alt_fill
        for cell in ws[row_index]:
            cell.font = Font(name="Calibri", size=11, color="243D55")
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.row_dimensions[row_index].height = 30

    for i, nome in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = LARGURAS[nome]
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 90

    if "Referência" in wb.sheetnames:
        ref = wb["Referência"]
        for row in range(2, ref.max_row):
            categoria = ref.cell(row, 1).value
            destino = primeira_linha.get(categoria)
            if destino:
                cell = ref.cell(row, 1)
                cell.hyperlink = f"#'Por Categoria'!A{destino}"
                cell.font = Font(name="Calibri", size=11, bold=False, color="0563C1", underline="single")

    wb.save(caminho_excel)
