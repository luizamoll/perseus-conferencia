"""Saídas operacionais e de consulta derivadas do mesmo resultado validado."""
import re
import textwrap
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo

from apresentacao import divergencias, resumo_resultado
from referencia import CATEGORIAS_REFERENCIA, referencia_visual, total_consolidado

MESES = ['Janeiro','Fevereiro','Marco','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']


def nome_excel(resultado):
    meta=resultado['metadados']
    m=re.fullmatch(r'(\d{4})/(0[1-9]|1[0-2])',meta.get('competencia_fatura',''))
    base=f'Resultado_Perseus_{MESES[int(m[2])-1]}_{m[1]}' if m else 'Resultado_Perseus'
    execucao=meta.get('execucao','')
    if execucao:
        if not re.fullmatch(r'[a-zA-Z0-9_-]+',execucao):raise ValueError('Identificação de execução inválida.')
        base+='_'+execucao
    return base+'.xlsx'


def nome_pdf(resultado):
    return nome_excel(resultado).replace('Resultado_', 'Relatorio_').replace('.xlsx','.pdf')


def nome_excel_mestre(resultado):
    competencia=(resultado.get('classificacao',{}) or {}).get('competencia_operacional','')
    if not competencia:
        meta=resultado.get('metadados',{})
        m=re.fullmatch(r'(\\d{4})/(0[1-9]|1[0-2])',meta.get('competencia_fatura',''))
        competencia=f'Fatura_Final_Ref_{MESES[int(m[2])-1]}_{m[1]}' if m else 'Fatura_Final_Ref'
    seguro=re.sub(r'[^A-Za-z0-9_-]+','_',competencia).strip('_')
    return f'Para_Colar_{seguro}.xlsx'


PLANILHA_MESTRE = [
    'Aluno','Curso','Data da Coleta','Status na Data da Coleta',
    'Status Atual','Unidade','Tipo de Curso','STATUS_CURSO','COMPETENCIA'
]


ALUNOS = [
    'Aluno','Identificador','Curso','Data da Coleta','Status na Data da Coleta',
    'Status Atual','Unidade','Tipo de Curso','STATUS_CURSO','COMPETENCIA',
    'Classificação','Motivo da revisão','Período de coleta','Competência da fatura',
    'Página','Tabela','Linha'
]
REVISAO = ['Aluno','Identificador','Curso','Problema encontrado','Motivo compreensível','Localização/origem','Período de coleta']
WIDTHS = {
    'Aluno':54,'Identificador':18,'Curso':38,'Data da Coleta':62,
    'Status na Data da Coleta':28,'Status Atual':26,'Unidade':20,'Tipo de Curso':20,
    'STATUS_CURSO':27,'COMPETENCIA':34,'Classificação':22,'Motivo da revisão':76,
    'Período de coleta':22,'Competência da fatura':24,'Página':10,'Tabela':10,'Linha':10,
    'Problema encontrado':28,'Motivo compreensível':84,'Localização/origem':36,
    'Indicador':44,'Resultado':88,'Categoria':38,'Quantidade':18,
}


def celula(cell, value, campo):
    if isinstance(value,str) and campo in ('Data da Coleta','Início','Fim') and re.fullmatch(r'\d{2}/\d{2}/\d{4}',value):
        try:value=datetime.strptime(value,'%d/%m/%Y')
        except ValueError:pass
    cell.value=value
    if isinstance(value,str):
        cell.data_type='s';cell.number_format='@'
    elif isinstance(value,datetime):cell.number_format='dd/mm/yyyy'
    elif isinstance(value,(int,float)):cell.number_format='#,##0'


def formatar(ws, campos):
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='B2' if ws.title=='Alunos' else 'A2'
    ws.auto_filter.ref=ws.dimensions
    ws.print_title_rows='1:1'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    if ws.title in ('Resumo','Referência'):ws.page_setup.fitToHeight=1
    ws.sheet_view.zoomScale=90
    for i,campo in enumerate(campos,1):ws.column_dimensions[get_column_letter(i)].width=WIDTHS[campo]
    for row in ws:
        height=30
        for c in row:
            header=c.row==1; campo=campos[c.column-1];width=WIDTHS[campo]
            value=c.value
            text=value.strftime('%d/%m/%Y') if isinstance(value,datetime) else str(value if value is not None else '')
            limit=max(10,int(width/1.15)-2)
            wrap=campo!='Identificador' and (len(text)>limit or '\n' in text)
            linecount=sum(max(1,len(textwrap.wrap(t,width=limit))) for t in text.split('\n')) if wrap else 1
            height=max(height,linecount*15+10)
            c.font=Font(name='Calibri',size=11,bold=header,color='FFFFFF' if header else '243D55')
            c.fill=PatternFill('solid',fgColor='243D55' if header else 'F0F4F7' if c.row%2==0 else 'FFFFFF')
            c.alignment=Alignment(vertical='center',horizontal='left' if header else 'center' if campo in ('Página','Tabela','Linha') else 'right' if isinstance(value,(int,float)) else 'left',wrap_text=wrap)
        ws.row_dimensions[row[0].row].height=height
    ws.print_options.horizontalCentered=True


def _texto_contraste(hexcor):
    cor=hexcor[-6:]
    r,g,b=(int(cor[i:i+2],16) for i in (0,2,4))
    return 'FFFFFF' if (0.299*r+0.587*g+0.114*b)<145 else '243D55'


def _criar_referencia(wb, resultado):
    classificacao=resultado.get('classificacao',{})
    consolidacao=classificacao.get('consolidacao',[])
    referencia=classificacao.get('referencia') or referencia_visual(resultado['metadados'].get('competencia_fatura',''))
    cor=referencia.get('cor','E7EFEA')
    texto=_texto_contraste(cor)
    ws=wb.create_sheet('Referência')
    ws.append([referencia.get('titulo','Referência'), 'Quantidade'])
    contagem={r.get('categoria'):int(r.get('quantidade',0)) for r in consolidacao}
    for categoria in CATEGORIAS_REFERENCIA:
        ws.append([categoria,contagem.get(categoria,0)])
    ws.append(['Total',total_consolidado(consolidacao)])
    formatar(ws,['Categoria','Quantidade'])
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor=cor);c.font=Font(name='Calibri',size=12,bold=True,color=texto)
    for c in ws[ws.max_row]:
        c.fill=PatternFill('solid',fgColor=cor);c.font=Font(name='Calibri',size=12,bold=True,color=texto)
    ws['A1']=referencia.get('titulo','Referência')
    ws.sheet_view.zoomScale=100
    return ws


def salvar_excel(resultado,caminho):
    resumos=resumo_resultado(resultado)
    wb=Workbook();wb.remove(wb.active)
    ws=wb.create_sheet('Resumo');ws.append(['Indicador','Resultado'])
    meta=resultado['metadados']
    classificacao=resultado.get('classificacao',{})
    linhas=[('Documento processado',meta['arquivo'])]
    for resumo in resumos:
        for key in ['Competência da fatura','Período de coleta','Início','Fim','Status geral','Total informado','Total encontrado','Diferença',
                    'Identificadores distintos','Registros conferidos (estrutura)','Classificados automaticamente','Registros para revisão','Alertas sem registro associado']:
            linhas.append((key,resumo[key]))
    linhas += [
        ('COMPETENCIA',classificacao.get('competencia_operacional','')),
        ('Fonte do STATUS_CURSO',(classificacao.get('fontes') or [''])[0]),
        ('Como interpretar','STATUS_CURSO segue a legenda oficial com a biblioteca operacional de divergências. Siglas ainda sem mapeamento permanecem visíveis para revisão.'),
        ('Escopo',meta.get('escopo','Primeiro bloco contíguo do relatório')),
        ('Natureza da planilha','Retrato editável desta execução. Alterar células não refaz a conferência do PDF.'),
        ('Execução',meta.get('execucao','Sem identificador de execução')),
        ('Processado em (UTC)',meta.get('processado_em','Não registrado')),
        ('SHA-256 do PDF',meta['sha256'])
    ]
    for i,(k,v) in enumerate(linhas,2):
        celula(ws.cell(i,1),k,'Indicador');celula(ws.cell(i,2),v,k)
    formatar(ws,['Indicador','Resultado'])
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ('Status geral','Total informado','Total encontrado','Diferença','Registros para revisão'):
            row[1].font=Font(name='Calibri',size=12,bold=True,color='243D55')
        if row[0].value in ('Diferença','Registros para revisão','Alertas sem registro associado'):
            row[1].fill=PatternFill('solid',fgColor='FFF0D8' if row[1].value else 'E6F2EC')

    _criar_referencia(wb, resultado)

    ws=wb.create_sheet('Alunos');ws.append(ALUNOS)
    audits=classificacao.get('auditoria') or []
    reviews={(d['Página'],d['Tabela'],d['Linha']):d['Motivo compreensível'] for d in divergencias(resultado) if d['_registro']}
    competencia=classificacao.get('competencia_operacional','')
    for i,r in enumerate(resultado['registros'],2):
        a=audits[i-2] if i-2<len(audits) else {}
        export={
            **r,
            'STATUS_CURSO':a.get('categoria_final') or 'NÃO LOCALIZADO',
            'COMPETENCIA':competencia,
            'Classificação':a.get('status_final','Não executada'),
            'Motivo da revisão':reviews.get((r['Página'],r['Tabela'],r['Linha']),'')
        }
        for col,k in enumerate(ALUNOS,1):celula(ws.cell(i,col),export.get(k,''),k)
    formatar(ws,ALUNOS)
    if ws.max_row>1:
        ws.conditional_formatting.add(
            f'K2:K{ws.max_row}',
            CellIsRule(operator='equal',formula=['"REVISAR"'],fill=PatternFill('solid',fgColor='FFF0D8'))
        )

    ws=wb.create_sheet('Divergências');ws.append(REVISAO)
    pendencias=divergencias(resultado)
    for i,r in enumerate(pendencias,2):
        for col,k in enumerate(REVISAO,1):celula(ws.cell(i,col),r.get(k,''),k)
    if not pendencias:
        celula(ws.cell(2,4),'Sem divergências identificadas','Problema encontrado')
        celula(ws.cell(2,5),'Não foram identificadas pendências de leitura, classificação ou quantidade nesta execução.','Motivo compreensível')
    formatar(ws,REVISAO)
    wb.save(caminho)


def salvar_excel_mestre(resultado,caminho):
    """Gera uma planilha enxuta, com a mesma estrutura usada na planilha mestre."""
    resumo_resultado(resultado)
    classificacao=resultado.get('classificacao',{})
    auditoria=classificacao.get('auditoria') or []
    registros=resultado.get('registros') or []
    if len(auditoria) != len(registros):
        raise ValueError('Não foi possível gerar a planilha mestre: classificação e registros estão desalinhados.')

    competencia=classificacao.get('competencia_operacional','')
    if not competencia:
        meta=resultado.get('metadados',{})
        m=re.fullmatch(r'(\\d{4})/(0[1-9]|1[0-2])',meta.get('competencia_fatura',''))
        competencia=f'Fatura_Final_Ref_{MESES[int(m[2])-1]}_{m[1]}' if m else 'Fatura_Final_Ref'

    wb=Workbook()
    ws=wb.active
    ws.title=competencia[:31]
    ws.sheet_view.showGridLines=True
    ws.freeze_panes='A2'
    ws.sheet_view.zoomScale=100
    ws.append(PLANILHA_MESTRE)

    for registro,classe in zip(registros,auditoria):
        linha={
            **registro,
            'STATUS_CURSO':classe.get('categoria_final') or 'NÃO LOCALIZADO',
            'COMPETENCIA':competencia,
        }
        ws.append([linha.get(campo,'') for campo in PLANILHA_MESTRE])

    larguras={
        'Aluno':52,'Curso':32,'Data da Coleta':24,'Status na Data da Coleta':28,
        'Status Atual':22,'Unidade':12,'Tipo de Curso':18,'STATUS_CURSO':24,'COMPETENCIA':34,
    }
    for i,campo in enumerate(PLANILHA_MESTRE,1):
        ws.column_dimensions[get_column_letter(i)].width=larguras[campo]
        cab=ws.cell(1,i)
        cab.font=Font(name='Calibri',size=11,bold=True,color='000000')
        cab.fill=PatternFill('solid',fgColor='70AD47')
        cab.alignment=Alignment(vertical='center',horizontal='center')

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='Calibri',size=11,color='000000')
            cell.alignment=Alignment(vertical='center',horizontal='left',wrap_text=False)
        data=row[2]
        if isinstance(data.value,str) and re.fullmatch(r'\\d{2}/\\d{2}/\\d{4}',data.value):
            try:
                data.value=datetime.strptime(data.value,'%d/%m/%Y')
                data.number_format='dd/mm/yyyy'
            except ValueError:
                pass

    if ws.max_row >= 2:
        tabela=Table(displayName='TabelaParaPlanilhaMestre',ref=f'A1:I{ws.max_row}')
        tabela.tableStyleInfo=TableStyleInfo(
            name='TableStyleMedium4',
            showFirstColumn=False,showLastColumn=False,
            showRowStripes=False,showColumnStripes=False,
        )
        ws.add_table(tabela)
    ws.auto_filter.ref=f'A1:I{ws.max_row}'
    wb.save(caminho)


def salvar_pdf(resultado,caminho):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle

    resumos=resumo_resultado(resultado);meta=resultado['metadados']
    classificacao=resultado.get('classificacao',{})
    referencia=classificacao.get('referencia') or referencia_visual(meta.get('competencia_fatura',''))
    consolidacao=classificacao.get('consolidacao',[])
    cor=colors.HexColor('#'+referencia.get('cor','E7EFEA'))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Cell',fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor('#243D55')))
    styles.add(ParagraphStyle(name='Head',parent=styles['Cell'],fontName='Helvetica-Bold',textColor=colors.white))
    def para(x,head=False):return Paragraph(escape(str(x if x is not None else '')).replace('\n','<br/>'),styles['Head' if head else 'Cell'])
    def table(rows,widths,header_color='#243D55'):
        t=Table([[para(v,i==0) for v in row] for i,row in enumerate(rows)],colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor(header_color)),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F0F4F7'),colors.white]),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)
        ]))
        return t

    doc=SimpleDocTemplate(str(caminho),pagesize=landscape(A4),leftMargin=32,rightMargin=32,topMargin=28,bottomMargin=36,title='Perseus - Relatório de conferência',author='Perseus')
    story=[Paragraph('Relatório de conferência',styles['Title']),para('Documento: '+meta['arquivo']),Spacer(1,10)]
    for r in resumos:
        story += [
            Paragraph('Fatura '+r['Competência da fatura']+' | Coleta '+r['Período de coleta'],styles['Heading2']),
            para('Intervalo da coleta: '+r['Início']+' a '+r['Fim']),Spacer(1,8),
            table([['Total informado','Total encontrado','Diferença','Registros para revisão'],[r[k] for k in ['Total informado','Total encontrado','Diferença','Registros para revisão']]],[194]*4),
            Spacer(1,8),Paragraph(r['Status geral'],styles['Heading3']),
            para(f"Identificadores distintos: {r['Identificadores distintos']} | Sem alertas de estrutura: {r['Registros conferidos (estrutura)']} | Classificados automaticamente: {r['Classificados automaticamente']} | Alertas sem registro associado: {r['Alertas sem registro associado']}")
        ]

    if consolidacao:
        story += [Spacer(1,14),Paragraph(referencia.get('titulo','Referência'),styles['Heading2'])]
        rows=[['STATUS_CURSO','Quantidade']]+[[item['categoria'],item['quantidade']] for item in consolidacao]+[['Total',total_consolidado(consolidacao)]]
        t=Table([[para(v,False) for v in row] for row in rows],colWidths=[310,120],hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),cor),('BACKGROUND',(0,-1),(-1,-1),cor),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTNAME',(0,-1),(-1,-1),'Helvetica-Bold'),
            ('TEXTCOLOR',(0,0),(-1,-1),colors.HexColor('#243D55')),
            ('ROWBACKGROUNDS',(0,1),(-1,-2),[colors.white,colors.HexColor('#F7F8F5')]),
            ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)
        ]))
        story += [t,para('COMPETENCIA: '+classificacao.get('competencia_operacional',''))]

    story += [Spacer(1,10),para('Diferença = encontrado menos informado. Conferência de quantidade não equivale a aprovação de regras de negócio.'),Paragraph('Divergências e pendências',styles['Heading2'])]
    pendencias=divergencias(resultado)
    for d in pendencias:
        if d['Origem']=='Conferência':story.extend([para(d['Motivo compreensível']),Spacer(1,8)])
    rows=[d for d in pendencias if d['Origem']!='Conferência']
    if rows:
        campos=['Aluno','Identificador','Curso','Motivo compreensível','Localização/origem']
        story.append(table([campos]+[[r[k] for k in campos] for r in rows],[160,72,108,318,118]))
    elif not pendencias:
        story.append(para('Sem divergências identificadas nesta execução.'))

    story += [
        Spacer(1,10),
        para('Escopo: '+meta.get('escopo','Primeiro bloco contíguo do relatório')),
        para('STATUS_CURSO calculado pela legenda oficial + biblioteca operacional; siglas ainda sem mapeamento permanecem para revisão.'),
        para('Execução: '+meta.get('execucao','Não registrada')+' | Processado em (UTC): '+meta.get('processado_em','Não registrado')),
        para('SHA-256 do PDF: '+meta['sha256'])
    ]
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#526577'))
        canvas.drawString(32,18,'Perseus | Fatura '+meta['competencia_fatura']+' | '+meta.get('execucao',''))
        canvas.drawRightString(810,18,'Página '+str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
