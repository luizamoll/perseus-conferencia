import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import app
import test_mvp as fixtures


class ApplicationTests(unittest.TestCase):
    def test_success_exports_zip_and_result(self):
        r=fixtures.ExtractionTests().run_pdf([fixtures.Page([fixtures.ROW])])
        with tempfile.TemporaryDirectory() as d:
            job={'arquivo':'origem.pdf'}
            with patch('app.extrair',return_value=r):
                app.processar(job,Path(d)/'entrada.pdf',Path(d))
            self.assertEqual(job['estado'],'concluido')
            self.assertEqual(job['conferencia'][0]['Diferença de linhas'],-1)
            self.assertTrue(job['referencia']['titulo'].startswith('Referência '))
            self.assertEqual(sum(x['quantidade'] for x in job['consolidacao']),len(r['registros']))
            with zipfile.ZipFile(Path(d)/'resultado.zip') as z:
                self.assertIn(job['excel'],z.namelist())
                self.assertIn(job['pdf'],z.namelist())
                self.assertIn(job['pdf_alunos'],z.namelist())
                self.assertTrue(z.read(job['pdf']).startswith(b'%PDF-'))
                self.assertTrue(z.read(job['pdf_alunos']).startswith(b'%PDF-'))
                self.assertIn('resultado.json',z.namelist())
                self.assertIn('biblioteca_divergencias.json',z.namelist())
            self.assertEqual(job['metadados']['arquivo'],'origem.pdf')

    def test_todas_pendencias_disponiveis_para_interface(self):
        rows=[]
        for i in range(15):
            row=list(fixtures.ROW);row[0]=f'Pessoa fictícia({i:04d})';row[1]='SETRA';rows.append(row)
        r=fixtures.ExtractionTests().run_pdf([fixtures.Page(rows)])
        with tempfile.TemporaryDirectory() as d:
            job={'arquivo':'teste.pdf'}
            with patch('app.extrair',return_value=r):app.processar(job,Path(d)/'entrada.pdf',Path(d))
            self.assertEqual(job['estado'],'concluido')
            self.assertGreater(job['total_pendencias'],10)
            self.assertEqual(len(job['divergencias']),job['total_pendencias'])
            self.assertEqual(job['divergencias'],app.divergencias(r))
            nao_localizado=next(x for x in job['consolidacao'] if x['categoria']=='NÃO LOCALIZADO')
            self.assertEqual(nao_localizado['quantidade'],15)

    def test_failure_state_is_actionable(self):
        job={}
        with patch('app.extrair',side_effect=ValueError('Cabeçalho Perseus não reconhecido.')):
            app.processar(job,Path('missing'),Path('unused'))
        self.assertEqual(job['estado'],'erro')
        self.assertIn('Cabeçalho',job['erro'])

    def test_download_xlsx_mime_and_whitelist(self):
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as d:
            dest=Path(d)/'test'/'resultado';dest.mkdir(parents=True)
            (dest/'alunos.xlsx').write_bytes(b'test-xlsx')
            (dest/'relatorio.pdf').write_bytes(b'%PDF-test')
            (dest/'lista-alunos.pdf').write_bytes(b'%PDF-alunos')
            handler=object.__new__(app.Pagina)
            handler.headers={'Host':'127.0.0.1:8876'}
            handler.server=SimpleNamespace(server_port=8876,destino=Path(d))
            responses=[];handler.responder=lambda *args:responses.append(args)
            with patch.dict(app.TRABALHOS,{'test':{'estado':'concluido','pdf':'relatorio.pdf','pdf_alunos':'lista-alunos.pdf'}}):
                handler.path='/download/test/alunos.xlsx';handler.do_GET();self.assertEqual(responses[-1][0],200);self.assertIn('spreadsheetml',responses[-1][2])
                handler.path='/download/test/relatorio.pdf';handler.do_GET();self.assertEqual(responses[-1][0],200);self.assertEqual(responses[-1][2],'application/pdf')
                handler.path='/download/test/lista-alunos.pdf';handler.do_GET();self.assertEqual(responses[-1][0],200);self.assertEqual(responses[-1][2],'application/pdf')
                handler.path='/download/test/entrada.pdf';handler.do_GET();self.assertEqual(responses[-1][0],404)

    def test_classificacao_no_excel_sem_alterar_registros(self):
        import copy
        from openpyxl import load_workbook
        curso_rows=[]
        for curso in ['DIR','CEXT, CHU, GN','CIB','SETRA']:
            row=list(fixtures.ROW);row[1]=curso;curso_rows.append(row)
        resultado=fixtures.ExtractionTests().run_pdf([fixtures.Page(curso_rows)])
        antes=copy.deepcopy(resultado['registros'])
        with tempfile.TemporaryDirectory() as d:
            job={'arquivo':'teste.pdf'}
            with patch('app.extrair',return_value=resultado):app.processar(job,Path(d)/'entrada.pdf',Path(d))
            self.assertEqual(job['estado'],'concluido')
            self.assertEqual(job['revisoes_classificacao'],1)
            self.assertEqual(resultado['registros'],antes)
            wb=load_workbook(Path(d)/'resultado'/job['excel'])
            self.assertEqual(wb.sheetnames,['Resumo','Referência','Alunos','Divergências'])
            rows=list(wb['Alunos'].values);headers=rows[0];data=[dict(zip(headers,row)) for row in rows[1:]]
            self.assertEqual([r['STATUS_CURSO'] for r in data],['GRADUAÇÃO','DI','GRADUAÇÃO','NÃO LOCALIZADO'])
            self.assertEqual([r['Classificação'] for r in data],['CLASSIFICADO','CLASSIFICADO','CLASSIFICADO','REVISAR'])
            self.assertTrue(data[3]['Motivo da revisão'])
            self.assertTrue(all(r['COMPETENCIA'] for r in data))
            ref=dict(wb['Referência'].values)
            self.assertEqual(ref['Total'],4)
            self.assertIn('Aluno',list(wb['Divergências'].values)[0])
            self.assertEqual(wb['Alunos'].freeze_panes,'B2')
            saved=json.loads((Path(d)/'resultado/resultado.json').read_text())
            self.assertEqual(saved['registros'],antes)
            self.assertEqual(len(saved['classificacao']['auditoria']),4)
            biblioteca=json.loads((Path(d)/'resultado/biblioteca_divergencias.json').read_text())
            self.assertEqual(biblioteca['siglas_nao_localizadas'],[{'sigla':'SETRA','ocorrencias':1}])

    def test_excel_recusa_classificacao_desalinhada(self):
        from classificacao import classificar_resultado
        from exportacao import salvar_excel
        r=fixtures.ExtractionTests().run_pdf([fixtures.Page([fixtures.ROW])])
        r['classificacao']=classificar_resultado(r)
        r['classificacao']['auditoria'][0]['registro_original']['Curso']='MEST'
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaisesRegex(ValueError,'não corresponde'):
                salvar_excel(r,Path(d)/'bad.xlsx')


if __name__=='__main__':
    unittest.main()
