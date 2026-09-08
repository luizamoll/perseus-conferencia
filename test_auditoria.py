import copy
import io
import tempfile
import unittest
from pathlib import Path
from datetime import datetime
from unittest.mock import patch
from openpyxl import load_workbook
from pypdf import PdfReader
import test_mvp as fixtures
from test_mvp import Page, ROW, BLOCK, CAMPOS
from perseus import validar_linha,salvar
from classificacao import classificar_resultado
from apresentacao import resumo_resultado,divergencias,motivo_humano
from exportacao import nome_excel,nome_pdf,salvar_excel,salvar_pdf
import app


class AuditoriaTests(unittest.TestCase):
    def resultado(self,rows=None):
        r=fixtures.ExtractionTests().run_pdf([Page(rows or [ROW])]);r['classificacao']=classificar_resultado(r);return r

    def test_header_competence_conflict(self):
        class Wrong(Page):
            def extract_text(self):return super().extract_text().replace('Competência: 2026/08','Competência: 2026/09')
        with self.assertRaisesRegex(ValueError,'Cabeçalhos contraditórios'):
            fixtures.ExtractionTests().run_pdf([Page([ROW]),Wrong([ROW])])

    def test_empty_first_block_not_resumed_later(self):
        r=fixtures.ExtractionTests().run_pdf([Page([]),Page([ROW],BLOCK.replace('07/2026 (','06/2026 (')),Page([ROW])])
        self.assertEqual(r['registros'],[])

    def test_each_date_token_validated(self):
        for value in ['26/06/2026, erro','26/06/2026,','26/06/2026abc','26/06/2026, 31/02/2026']:
            row=list(ROW);row[2]=value
            _,problems=validar_linha(row,datetime(2026,6,26),datetime(2026,7,25))
            self.assertTrue(any('Data' in p or 'data' in p for p in problems))

    def test_alerts_have_exact_source_and_human_identity(self):
        row=list(ROW);row[3]=''
        r=self.resultado([row]);d=[x for x in divergencias(r) if x['_registro']][0]
        self.assertEqual((d['Identificador'],d['Tabela'],d['Linha']),('0001',1,3))
        self.assertIn('Tabela 1',d['Localização/origem']);self.assertTrue(d['Aluno'])

    def test_review_counts_unique_records_not_reasons(self):
        row=list(ROW);row[1]='SETRA';row[3]='';row[4]=''
        r=self.resultado([row]);summary=resumo_resultado(r)[0]
        self.assertEqual(summary['Registros para revisão'],1)
        self.assertGreater(len(r['revisao']),1)
        self.assertEqual(len([d for d in divergencias(r) if d['_registro']]),1)
        texto=' '.join(d['Motivo compreensível'] for d in divergencias(r))
        self.assertIn('não localizado',texto.lower())

    def test_total_difference_is_not_an_invented_student(self):
        r=self.resultado();ds=divergencias(r)
        self.assertEqual(ds[0]['Origem'],'Conferência');self.assertEqual(ds[0]['Identificador'],'')
        self.assertEqual(resumo_resultado(r)[0]['Registros para revisão'],0)
        self.assertEqual(resumo_resultado(r)[0]['Status geral'],'Divergência de quantidade')

    def test_zero_pending_explicit(self):
        row=list(ROW);row[0]='Outro(0002)';r=self.resultado([ROW,row])
        self.assertEqual(divergencias(r),[])
        with tempfile.TemporaryDirectory() as d:
            salvar_excel(r,Path(d)/'a.xlsx');salvar_pdf(r,Path(d)/'a.pdf')
            w=load_workbook(Path(d)/'a.xlsx');self.assertEqual(w['Divergências']['D2'].value,'Sem divergências identificadas')
            self.assertIn('Sem divergências', ''.join(p.extract_text() for p in PdfReader(Path(d)/'a.pdf').pages))

    def test_stale_summary_or_competence_is_rejected(self):
        for change in ['count','competence','classification']:
            r=self.resultado()
            if change=='count':r['conferencia'][0]['Linhas extraídas']+=1
            elif change=='competence':r['registros'][0]['Competência da fatura']='2026/09'
            else:r['classificacao']['auditoria'][0]['registro_original']['Curso']='OTHER'
            with self.assertRaises(ValueError):resumo_resultado(r)

    def test_identifiers_exact_text_and_safe_strings(self):
        for ident in ['0000123','DIs0001','12345678901234567890','=1+1']:
            row=list(ROW);row[0]='Nome Completo('+ident+')';r=self.resultado([row]);old=copy.deepcopy(r)
            with tempfile.TemporaryDirectory() as d:
                salvar_excel(r,Path(d)/'x.xlsx');w=load_workbook(Path(d)/'x.xlsx')
                self.assertEqual(w['Alunos']['B2'].value,ident);self.assertEqual(w['Alunos']['B2'].data_type,'s');self.assertEqual(w['Alunos']['B2'].number_format,'@')
            self.assertEqual(r,old)

    def test_single_date_native_multiple_preserved(self):
        row=list(ROW);row[2]='26/06/2026';r=self.resultado([row])
        with tempfile.TemporaryDirectory() as d:
            salvar_excel(r,Path(d)/'x.xlsx');w=load_workbook(Path(d)/'x.xlsx')
            self.assertEqual(w['Alunos']['D2'].value,datetime(2026,6,26));self.assertEqual(w['Alunos']['D2'].number_format,'dd/mm/yyyy')

    def test_export_preserves_all_primary_fields(self):
        r=self.resultado();old=copy.deepcopy(r)
        with tempfile.TemporaryDirectory() as d:
            salvar(r,Path(d)/'out');w=load_workbook(Path(d)/'out'/nome_excel(r));headers=[c.value for c in w['Alunos'][1]]
            values=dict(zip(headers,[c.value for c in w['Alunos'][2]]))
            for k in CAMPOS:self.assertEqual(values[k],r['registros'][0][k])
            self.assertEqual(w.sheetnames,['Resumo','Referência','Alunos','Divergências'])
            self.assertEqual(values['STATUS_CURSO'],'MÚLTIPLAS CATEGORIAS')
            self.assertTrue(values['COMPETENCIA'])
            self.assertTrue(all(ws.freeze_panes and ws.auto_filter.ref and not ws.merged_cells for ws in w))
            self.assertTrue((Path(d)/'out'/nome_pdf(r)).exists())
        self.assertEqual(r,old)

    def test_different_runs_have_different_download_names(self):
        names=[]
        for i in ['run_1','run_2']:
            with tempfile.TemporaryDirectory() as d:
                r=self.resultado();job={'id':i,'arquivo':'Mesmo.pdf'}
                with patch('app.extrair',return_value=r):app.processar(job,Path(d)/'in.pdf',Path(d))
                self.assertEqual(job['estado'],'concluido');self.assertEqual(job['resumo'],resumo_resultado(r))
                self.assertIn(i,job['excel']);self.assertIn(i,job['pdf']);names.append(job['excel'])
        self.assertNotEqual(*names)

    def test_download_no_cache_and_attachment(self):
        h=object.__new__(app.Pagina);h.wfile=io.BytesIO();headers={}
        h.send_response=lambda status:None;h.send_header=lambda k,v:headers.update({k:v});h.end_headers=lambda:None
        h.responder(200,b'%PDF-test','application/pdf','resultado.pdf')
        self.assertEqual(headers['Cache-Control'],'no-store');self.assertIn('attachment;',headers['Content-Disposition'])

    def test_biblioteca_operacional_separa_confirmado_de_historico(self):
        from classificacao import classificar_registro
        confirmados={'TEOC':'GRADUAÇÃO','LPI':'GRADUAÇÃO','EFB':'GRADUAÇÃO','PCG':'POS'}
        for course,categoria in confirmados.items():
            a=classificar_registro({'Curso':course});self.assertEqual(a['categoria_final'],categoria);self.assertEqual(a['status_final'],'CLASSIFICADO')
        historicos={
            'CIB':'GRADUAÇÃO','GEQUA':'POS','ESGSC, PSO':'GRADUAÇÃO',
            'DIREM, MBAAD':'GRADUAÇÃO','DA, DIR':'GRADUAÇÃO'
        }
        for course,categoria in historicos.items():
            a=classificar_registro({'Curso':course},permitir_historico=True);self.assertEqual(a['categoria_final'],categoria);self.assertEqual(a['status_final'],'CLASSIFICADO')
        self.assertEqual(classificar_registro({'Curso':'CIB'})['categoria_final'],'DI')
        a=classificar_registro({'Curso':'SETRA'});self.assertEqual(a['categoria_final'],'NÃO LOCALIZADO');self.assertEqual(a['status_final'],'REVISAR')
        self.assertTrue(all(motivo_humano(m) for m in a['motivos_revisao']))
        a=classificar_registro({'Curso':'15, CEXT, CSAP, GN'})
        self.assertEqual(a['categoria_final'],'DI');self.assertEqual(a['status_final'],'REVISAR')


if __name__=='__main__':
    unittest.main()
