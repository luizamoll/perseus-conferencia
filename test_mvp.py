import copy
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from perseus import CAMPOS, extrair, salvar

HEADER='Contrato: 123 Competência: 2026/08 Quantidade Faturada: 2'
BLOCK='07/2026 (26/06/2026 até 25/07/2026) Coletado: 2 Calculado: 2'
ROW=['Exemplo(0001)', 'DIR, EVE', '26/06/2026, 26/06/2026', 'MAT, MAT', 'MAT, MAT', '1, 1', 'S, S']
class Page:
    def __init__(self, rows, block=BLOCK): self.rows,self.block=rows,block
    def extract_text(self): return HEADER+'\n'+self.block
    def extract_tables(self): return [[[self.block], CAMPOS, *self.rows]]
    def close(self): pass
class PDF:
    def __init__(self, pages): self.pages=pages
    def __enter__(self): return self
    def __exit__(self,*a): pass

class ExtractionTests(unittest.TestCase):
    def run_pdf(self,pages):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'source.pdf'; p.write_bytes(b'%PDF-test')
            with patch('perseus.pdfplumber.open', return_value=PDF(pages)):
                return extrair(p)
    def test_duplicates_preserved_and_flagged(self):
        r=self.run_pdf([Page([ROW,ROW])])
        self.assertEqual(len(r['registros']),2)
        self.assertEqual(r['conferencia'][0]['Identificadores distintos'],1)
        self.assertEqual(len(r['revisao']),2)
    def test_empty_table_still_reconciled(self):
        r=self.run_pdf([Page([])])
        self.assertEqual(r['conferencia'][0]['Diferença de linhas'],-2)
        self.assertTrue(r['revisao'])
    def test_malformed_row_is_retained_in_audit(self):
        r=self.run_pdf([Page([ROW,['broken']])])
        self.assertEqual(len(r['registros']),1)
        self.assertEqual(r['revisao'][0]['dados_originais'],['broken'])
    def test_contradictory_totals_raise(self):
        with self.assertRaisesRegex(ValueError,'contraditórios'):
            self.run_pdf([Page([ROW]),Page([ROW],BLOCK.replace('Coletado: 2','Coletado: 3'))])
    def test_next_period_excluded(self):
        r=self.run_pdf([Page([ROW]),Page([ROW],BLOCK.replace('07/2026 (','06/2026 ('))])
        self.assertEqual(len(r['registros']),1)
    def test_excel_roundtrip_and_technical_exports(self):
        r=self.run_pdf([Page([ROW])])
        r['registros'][0]['Aluno']='=HYPERLINK("bad")'
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'result';salvar(r,out)
            wb=load_workbook(next(out.glob('*.xlsx')))
            self.assertEqual(wb.sheetnames,['Resumo','Alunos','Divergências'])
            ws=wb['Alunos'];self.assertEqual(ws['A2'].data_type,'s')
            self.assertEqual(ws['B2'].value,'0001')
            self.assertEqual(ws['B2'].data_type,'s')
            self.assertEqual(ws['F2'].value,ROW[2])
            self.assertEqual(ws.freeze_panes,'B2')
            self.assertTrue(ws.auto_filter.ref)
            self.assertEqual(json.loads((out/'resultado.json').read_text()),r)
            self.assertIn("'=HYPERLINK",(out/'alunos.csv').read_text(encoding='utf-8-sig'))

@unittest.skipUnless(os.getenv('PERSEUS_AGOSTO') and os.getenv('PERSEUS_SETEMBRO'), 'Defina PERSEUS_AGOSTO e PERSEUS_SETEMBRO para testes reais')
class RealPDFTests(unittest.TestCase):
    def test_real_pdfs(self):
        for mes, total, extraidos, periodo in [('AGOSTO',2050,2050,'07/2026'),('SETEMBRO',2213,2211,'08/2026')]:
            with self.subTest(mes=mes):
                p=Path(os.environ['PERSEUS_'+mes]); before=hashlib.sha256(p.read_bytes()).hexdigest()
                r=extrair(p)
                c=r['conferencia'][0]
                self.assertEqual(c['Coletado informado'],total)
                self.assertEqual(c['Linhas extraídas'],extraidos)
                self.assertEqual(c['Identificadores distintos'],extraidos)
                self.assertEqual(c['Período de coleta'],periodo)
                self.assertEqual(c['Diferença de linhas'],extraidos-total)
                self.assertEqual(r['revisao'],[])
                self.assertTrue(all(all(row[k] for k in CAMPOS) for row in r['registros']))
                self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),before)
                baseline=os.getenv('PERSEUS_BASELINE')
                if baseline:
                    old=json.loads((Path(baseline)/('baseline_'+mes.lower())/'resultado.json').read_text())
                    self.assertEqual(r['registros'],old['registros'])
                    self.assertEqual(r['conferencia'],old['conferencia'])
                out=os.getenv('PERSEUS_TEST_OUTPUT')
                if out: salvar(r,Path(out)/mes.lower())
