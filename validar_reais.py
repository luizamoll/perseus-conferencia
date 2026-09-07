"""Validação reproduzível dos dois exemplos: processamento real ou resultados já gerados."""
import argparse
import hashlib
import io
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from app import processar
from apresentacao import resumo_resultado,divergencias
from perseus import BLOCO,CABECALHO,CAMPOS

BASE=Path(__file__).resolve().parent


def validar(pdf,resultado,excel,relatorio):
    r=json.loads(Path(resultado).read_text(encoding='utf-8'))
    source_hash=hashlib.sha256(Path(pdf).read_bytes()).hexdigest()
    assert source_hash==r['metadados']['sha256']
    summary=resumo_resultado(r)[0]
    text_reader=PdfReader(pdf)
    first_text=text_reader.pages[0].extract_text()
    header=CABECALHO.search(first_text);block=BLOCO.search(first_text)
    assert header and block
    assert (summary['Competência da fatura'],summary['Período de coleta'],summary['Total informado'])==(header[2],block[1],int(block[4]))
    text_ids=[]
    for page in text_reader.pages:
        text=page.extract_text();blocks=list(BLOCO.finditer(text))
        if blocks and all(b[1]!=block[1] for b in blocks):break
        text_ids.extend(re.findall(r'\((DIs\d+|\d+)\)',text))
    assert Counter(text_ids)==Counter(x['Identificador'] for x in r['registros'])
    wb=load_workbook(excel)
    assert wb.sheetnames==['Resumo','Alunos','Divergências']
    saved_summary=dict(wb['Resumo'].values)
    for k,v in summary.items():
        if k in ('Início','Fim'):v=datetime.strptime(v,'%d/%m/%Y')
        assert saved_summary[k]==v,(k,saved_summary[k],v)
    sheet=wb['Alunos'];headers=[c.value for c in sheet[1]]
    assert sheet.max_row==len(r['registros'])+1
    for actual,cells in zip(r['registros'],sheet.iter_rows(min_row=2)):
        saved=dict(zip(headers,[c.value for c in cells]))
        for k in CAMPOS+['Identificador','Página','Tabela','Linha']:
            v=saved[k];v=v.strftime('%d/%m/%Y') if isinstance(v,datetime) else v
            assert v==actual[k],(k,v,actual[k])
        assert cells[1].data_type=='s' and cells[1].number_format=='@'
    assert all(w.freeze_panes and w.auto_filter.ref and not w.merged_cells for w in wb)
    assert max(d.height for w in wb for d in w.row_dimensions.values())<=100
    assert wb['Divergências'].max_row==max(1,len(divergencias(r)))+1
    report=PdfReader(relatorio);text=''.join(p.extract_text() for p in report.pages)
    normal=lambda x:re.sub(r'\s+','',x)
    for v in [header[2],block[1],str(summary['Total informado']),str(summary['Total encontrado']),summary['Status geral']]:assert v in text,v
    for d in divergencias(r):
        assert normal(d['Motivo compreensível']) in normal(text)
        if d['Identificador']:assert d['Identificador'] in text
    return {'resumo':summary,'paginas_pdf_saida':len(report.pages),'identificadores_texto_independente':len(text_ids),'sha256':source_hash}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--resultados-existentes',type=Path,help='Pasta com agosto/resultado.json e setembro/resultado.json, XLSX e PDF respectivos.')
    parser.add_argument('--saida',type=Path,help='Pasta nova para processar os dois exemplos quando resultados existentes não são fornecidos.')
    args=parser.parse_args();res={}
    if not args.resultados_existentes and not args.saida:parser.error('Informe --saida ou --resultados-existentes.')
    for mes in ('Agosto','Setembro'):
        source=BASE/'exemplos'/f'{mes}.pdf'
        if args.resultados_existentes:
            folder=args.resultados_existentes/mes.lower()
        else:
            folder=args.saida/mes.lower();folder.mkdir(parents=True,exist_ok=False)
            job={'arquivo':source.name};processar(job,source,folder)
            assert job['estado']=='concluido',job
            folder=folder/'resultado'
        res[mes]=validar(source,folder/'resultado.json',next(folder.glob('*.xlsx')),next(folder.glob('*.pdf')))
        print(mes,json.dumps(res[mes],ensure_ascii=False),flush=True)
    print('Dois PDFs reais, resumos, registros, identificadores e saídas: OK',flush=True)

if __name__=='__main__':main()
